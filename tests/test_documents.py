"""Tests for Documents: documents.py and how the server uses attached documents."""
import os, sys; sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), os.pardir)); import emvryx_paths  # noqa: E402,F401  (the code folders become importable)
import base64
import io
import json
import unittest
import zipfile

import app_engine as E
import documents as D
from tests.test_app_server import Base, Client


def make_docx(paragraphs, table=None):
    from docx import Document
    d = Document()
    for p in paragraphs:
        d.add_paragraph(p)
    if table:
        t = d.add_table(rows=len(table), cols=len(table[0]))
        for r, row in enumerate(table):
            for c, v in enumerate(row):
                t.cell(r, c).text = v
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def make_pdf(pages):
    from fpdf import FPDF
    pdf = FPDF()
    for text in pages:
        pdf.add_page()
        pdf.set_font("Helvetica", size=12)
        pdf.multi_cell(0, 8, text)
    return bytes(pdf.output())


LONG = "\n\n".join(f"Section {i}. " + ("Filler sentence about nothing in particular. " * 25) for i in range(40))
LONG += "\n\nThe treasure is buried under the old oak tree behind the lighthouse. " + ("More filler text for padding. " * 30)


class Extract(unittest.TestCase):
    def test_text_markdown_csv(self):
        self.assertEqual(D.extract_text("a.txt", b"hello\r\nworld"), "hello\nworld")
        self.assertIn("a,b", D.extract_text("t.csv", b"a,b\n1,2"))
        self.assertIn("café", D.extract_text("n.md", "café".encode("latin-1")))

    def test_html_loses_tags_and_scripts(self):
        t = D.extract_text("p.html", b"<html><script>evil()</script><p>Hello <b>there</b></p><p>Next &amp; more</p></html>")
        self.assertIn("Hello there", t)
        self.assertIn("Next & more", t)
        self.assertNotIn("evil", t)

    def test_docx_with_table(self):
        t = D.extract_text("r.docx", make_docx(["First paragraph.", "Second."], [["Name", "Score"], ["Ann", "9"]]))
        self.assertIn("First paragraph.", t)
        self.assertIn("Ann | 9", t)

    def test_pdf_has_page_marks(self):
        t = D.extract_text("r.pdf", make_pdf(["Alpha page text.", "Beta page text."]))
        self.assertIn("[Page 1]", t)
        self.assertIn("Beta page text.", t)

    def test_problems_are_plain_messages(self):
        cases = [("a.txt", b""), ("a.exe", b"MZ\x00\x01"), ("a.docx", b"not a docx"), ("a.pdf", b"not a pdf"), ("a.txt", b"x" * (D.MAX_FILE_BYTES + 1)),
                 ("a.bin", b"\x00\x01\x02\x03"), ("a.txt", b"   \n  ")]
        for name, data in cases:
            with self.assertRaises(ValueError, msg=name) as cm:
                D.extract_text(name, data)
            self.assertTrue(str(cm.exception))

    def test_a_pdf_without_text_is_explained(self):
        from fpdf import FPDF
        pdf = FPDF()
        pdf.add_page()
        with self.assertRaises(ValueError) as cm:
            D.extract_text("scan.pdf", bytes(pdf.output()))
        self.assertIn("no readable text", str(cm.exception))


class Library(unittest.TestCase):
    def test_add_list_get_delete(self):
        lib = D.Library()
        a = lib.add("notes.txt", b"alpha beta")
        b = lib.add("../evil/name.md", b"gamma")
        self.assertEqual([d["name"] for d in lib.list()], [lib.info(b["id"])["name"], "notes.txt"])
        self.assertNotIn("/", lib.info(b["id"])["name"])
        self.assertEqual(lib.text(a["id"]), "alpha beta")
        self.assertTrue(lib.delete(a["id"]))
        self.assertFalse(lib.delete(a["id"]))
        self.assertIsNone(lib.text(a["id"]))
        lib.clear()
        self.assertEqual(lib.list(), [])

    def test_limit_on_how_many(self):
        lib = D.Library()
        for i in range(D.MAX_DOCS):
            lib.add(f"d{i}.txt", b"x")
        with self.assertRaises(ValueError):
            lib.add("one-more.txt", b"x")


class Choosing(unittest.TestCase):
    def test_a_short_document_goes_in_whole(self):
        lib = D.Library()
        d = lib.add("s.txt", b"The meeting is on Tuesday at noon.")
        notes = D.notes_for(lib, [d["id"]], "when is the meeting", budget=8000)
        self.assertEqual(len(notes), 1)
        self.assertIn("Tuesday", notes[0]["text"])
        self.assertEqual(notes[0]["title"], "Document: s.txt")

    def test_a_long_document_is_searched(self):
        lib = D.Library()
        d = lib.add("long.txt", LONG.encode())
        notes = D.notes_for(lib, [d["id"]], "where is the treasure buried", budget=3000)
        self.assertLessEqual(sum(len(n["text"]) for n in notes), 4500)
        self.assertTrue(any("old oak tree" in n["text"] for n in notes))

    def test_overview_questions_are_spread_over_the_document(self):
        pieces = D.chunks(LONG)
        picked = D.pick(pieces, "Please summarize this document", 4000)
        self.assertEqual(picked[0], 0)
        self.assertGreater(picked[-1], len(pieces) // 2)

    def test_chunks_cover_everything_in_order(self):
        pieces = D.chunks(LONG, 1200)
        self.assertTrue(all(len(p) <= 1900 for p in pieces))
        self.assertIn("lighthouse", pieces[-1])

    def test_hint_names_the_document(self):
        self.assertIn('"plan.pdf"', D.hint(["plan.pdf"]))


class ChatStore(unittest.TestCase):
    def test_docs_per_chat_and_private(self):
        s = E.ChatStore()
        a, p = s.new_chat(), s.new_chat(private=True)
        s.set_docs(a, [3, 4])
        s.set_docs(p, [5])
        self.assertEqual((s.docs(a), s.docs(p), s.docs(s.new_chat())), ([3, 4], [5], []))


def b64(data):
    return base64.b64encode(data).decode()


class Server(Base):
    backend = None

    def factory(model, question):                                   # noqa: N805
        Server.backend = E.FakeBackend()
        return Server.backend

    factory = staticmethod(factory)

    def decode(self, ids):
        return self.state.tok_for({"id": "flare-2"}).decode(ids)

    def upload(self, who, name, data):
        return who.call("POST", "/api/documents", {"name": name, "data": b64(data)})

    def test_upload_list_delete(self):
        s, j = self.upload(self.sam, "plan.txt", b"The launch is on Friday.")
        self.assertEqual(s, 200)
        self.assertEqual(j["name"], "plan.txt")
        s, listing = self.sam.call("GET", "/api/documents")
        self.assertEqual([d["name"] for d in listing["documents"]], ["plan.txt"])
        self.assertEqual(self.owner.call("GET", "/api/documents")[1]["documents"], [])        # people don't see each other's documents
        self.assertEqual(self.owner.call("DELETE", f"/api/documents/{j['id']}")[1]["ok"], False)
        self.assertEqual(self.sam.call("DELETE", f"/api/documents/{j['id']}")[1]["ok"], True)

    def test_bad_files_are_refused_with_a_message(self):
        s, j = self.upload(self.sam, "a.exe", b"MZ\x00\x01")
        self.assertEqual(s, 400)
        self.assertIn("can't read", j["message"])
        self.assertEqual(self.sam.call("POST", "/api/documents", {"name": "a.txt", "data": "!!!notbase64!!!"})[0], 400)

    def test_needs_a_sign_in(self):
        self.assertEqual(Client(self.port).call("GET", "/api/documents")[0], 401)
        self.assertEqual(Client(self.port).call("POST", "/api/documents", {"name": "a.txt", "data": b64(b"x")})[0], 401)

    def test_an_attached_document_reaches_the_model_but_is_never_saved_in_the_chat(self):
        doc = self.upload(self.sam, "plan.txt", b"The launch is on Friday at nine.")[1]
        s, j = self.sam.call("POST", "/api/chats", {"docs": [doc["id"]]})
        cid = j["id"]
        self.ask(self.sam, cid, "When is the launch")
        prompt = self.decode(Server.backend.prompts[-1])
        self.assertIn("launch is on Friday at nine", prompt)
        self.assertIn("plan.txt", prompt)
        chat = self.sam.call("GET", f"/api/chats/{cid}")[1]
        self.assertEqual([d["name"] for d in chat["docs"]], ["plan.txt"])
        self.assertNotIn("Friday", chat["messages"][0]["content"])

    def test_attach_later_and_detach(self):
        doc = self.upload(self.sam, "plan.txt", b"The launch is on Friday.")[1]
        cid = self.new_chat(self.sam)
        self.ask(self.sam, cid, "hello")
        self.assertNotIn("launch is on", self.decode(Server.backend.prompts[-1]))
        s, j = self.sam.call("POST", f"/api/chats/{cid}/docs", {"ids": [doc["id"]]})
        self.assertEqual([d["name"] for d in j["docs"]], ["plan.txt"])
        self.ask(self.sam, cid, "when is it")
        self.assertIn("launch is on Friday", self.decode(Server.backend.prompts[-1]))
        self.sam.call("POST", f"/api/chats/{cid}/docs", {"ids": []})
        self.ask(self.sam, cid, "and now")
        self.assertNotIn("launch is on", self.decode(Server.backend.prompts[-1]))

    def test_someone_elses_document_cannot_be_attached(self):
        doc = self.upload(self.owner, "secret.txt", b"The code word is pelican.")[1]
        cid = self.new_chat(self.sam)
        j = self.sam.call("POST", f"/api/chats/{cid}/docs", {"ids": [doc["id"]]})[1]
        self.assertEqual(j["docs"], [])
        self.ask(self.sam, cid, "what is the code word")
        self.assertNotIn("pelican", self.decode(Server.backend.prompts[-1]))

    def test_export_includes_documents_and_delete_everything_clears_them(self):
        self.upload(self.sam, "plan.txt", b"The launch is on Friday.")
        s, _, raw = self.sam.call("GET", "/api/export", raw=True)
        z = zipfile.ZipFile(io.BytesIO(raw))
        self.assertEqual(json.loads(z.read("documents.json"))[0]["name"], "plan.txt")
        self.sam.call("POST", "/api/delete-everything", {})
        self.assertEqual(self.sam.call("GET", "/api/documents")[1]["documents"], [])


if __name__ == "__main__":
    unittest.main()
