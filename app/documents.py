"""
documents.py - Documents: files a person adds (PDF, Word, text ...) so a chat can answer questions about them.

WHAT THIS FILE DOES
    extract_text   a file's bytes -> plain text (PDF with [Page N] marks, Word with its tables, text/Markdown/CSV/JSON/code, HTML without tags)
    Library        one person's saved documents (SQLite, one file per person like the chats): add, list, get, delete, clear, with limits
    notes_for      the parts of the attached documents that matter for THIS question, as "notes" for the model

    Nothing is trained and nothing is sent anywhere: the text is kept on the computer that runs the app, in that person's own file, and the
    model only ever sees the few passages chosen for each question. A short document goes in whole. A long one is searched:
    the passages that share the most (rarer) words with the question win; for "summarize this" style questions the passages are spread over
    the whole document instead, so the answer is not about just one page.

    Limits (plain messages, never a crash): 10 MB per file, about 1.5 million letters of text per document, 40 documents per person.
"""
import html
import io
import math
import re
import sqlite3
import time

from safe_db import SafeDB

MAX_FILE_BYTES = 10_000_000
MAX_CHARS = 1_500_000
MAX_DOCS = 40
TEXT_EXTS = {".txt", ".md", ".markdown", ".csv", ".tsv", ".json", ".log", ".py", ".js", ".ts", ".java", ".c", ".cpp", ".h", ".cs", ".go", ".rs",
             ".rb", ".php", ".sql", ".yaml", ".yml", ".toml", ".ini", ".xml", ".css", ".sh", ".bat", ".tex", ".rst"}
HTML_EXTS = {".html", ".htm"}
IMAGE_EXTS = {".png", ".jpg", ".jpeg", ".webp", ".bmp", ".gif", ".heic"}
KINDS = "PDF, Word (.docx), HTML, a picture or screenshot (its text is read), or a text file (.txt, .md, .csv, .json, code)"
IMAGE_KINDS = {e.lstrip(".") for e in IMAGE_EXTS}
_ocr = None                                                                 # the text reader is loaded on first use (it takes a second)


def warm_ocr():
    """Load the picture reader ahead of time (the app does this in the background when it starts), so the first picture isn't slow. Quietly does nothing if it isn't installed."""
    global _ocr
    try:
        if _ocr is None:
            from rapidocr_onnxruntime import RapidOCR
            _ocr = RapidOCR()
    except Exception:
        pass


def ocr_image(data):
    """The text in a picture, read top to bottom. The models can't see pictures yet; this is what lets a screenshot of text be used."""
    global _ocr
    try:
        from PIL import Image
        if _ocr is None:
            from rapidocr_onnxruntime import RapidOCR
            _ocr = RapidOCR()
    except ImportError:
        raise ValueError("Reading text in pictures needs one extra package on the computer that runs the app: pip install rapidocr-onnxruntime") from None
    try:
        img = Image.open(io.BytesIO(data)).convert("RGB")
        if max(img.size) > 3000:
            img.thumbnail((3000, 3000))
        buf = io.BytesIO()
        img.save(buf, "PNG")
        result, _ = _ocr(buf.getvalue())
    except Exception:
        raise ValueError("I couldn't open that picture. Is it a PNG, JPG or WebP?") from None
    rows = sorted(result or [], key=lambda r: (round(min(p[1] for p in r[0]) / 12), min(p[0] for p in r[0])))   # lines, then left to right
    return "\n".join(r[1] for r in rows)


def _ext(name):
    m = re.search(r"\.[A-Za-z0-9]+$", name or "")
    return m.group(0).lower() if m else ""


def _clean(text):
    text = re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f\x7f]", "", text.replace("\r\n", "\n").replace("\r", "\n"))
    text = re.sub(r"[ \t]+\n", "\n", text)
    return re.sub(r"\n{3,}", "\n\n", text).strip()


def _decode(data):
    if data.startswith((b"\xff\xfe", b"\xfe\xff")):                         # only a byte-order mark says UTF-16 (otherwise old Western text would read as garbage)
        try:
            return data.decode("utf-16")
        except UnicodeError:
            pass
    try:
        return data.decode("utf-8-sig")
    except UnicodeDecodeError:
        return data.decode("cp1252", "replace")


PICTURE_ONLY = "[A picture with no readable text]"      # a photo with no words in it: accepted only when the picture reader can look at it


def extract_text(name, data, allow_picture=False):
    """Plain text of an uploaded file. A problem with the file is a ValueError with a message for the person."""
    if not data:
        raise ValueError("That file is empty.")
    if len(data) > MAX_FILE_BYTES:
        raise ValueError(f"That file is too big ({len(data) // 1_000_000} MB). The limit is {MAX_FILE_BYTES // 1_000_000} MB.")
    ext = _ext(name)
    if ext == ".pdf":
        try:
            from pypdf import PdfReader
        except ImportError:
            raise ValueError("Reading PDFs needs one extra package on the computer that runs the app: pip install pypdf") from None
        try:
            reader = PdfReader(io.BytesIO(data))
            if getattr(reader, "is_encrypted", False):
                if not reader.decrypt(""):
                    raise ValueError("This PDF is password protected. Remove the password and add it again.")
            pages = []
            for i, page in enumerate(reader.pages, 1):
                t = (page.extract_text() or "").strip()
                if t:
                    pages.append(f"[Page {i}]\n{t}")
        except ValueError:
            raise
        except Exception:
            raise ValueError("I couldn't read that PDF. It may be damaged.") from None
        text = "\n\n".join(pages)
        if not text.strip():
            raise ValueError("This PDF has no readable text (it may be a scan or pictures). Try a text version of it.")
    elif ext == ".docx":
        try:
            from docx import Document
        except ImportError:
            raise ValueError("Reading Word files needs one extra package on the computer that runs the app: pip install python-docx") from None
        try:
            d = Document(io.BytesIO(data))
        except Exception:
            raise ValueError("I couldn't read that Word file. Is it a .docx (not the older .doc)?") from None
        parts = [p.text for p in d.paragraphs if p.text.strip()]
        for t in d.tables:
            for row in t.rows:
                parts.append(" | ".join(c.text.strip() for c in row.cells))
        text = "\n".join(parts)
    elif ext in IMAGE_EXTS:
        text = ocr_image(data)
        if not text.strip() and allow_picture:                       # no words, but the picture reader can look at it when it is asked about
            return PICTURE_ONLY
        if not text.strip():
            raise ValueError("I couldn't find any text in that picture. Emvryx can read the words in a picture or screenshot, but it can't describe photos yet.")
        text = "[Text read from a picture]\n" + text
    elif ext in HTML_EXTS:
        raw = _decode(data)
        raw = re.sub(r"(?is)<(script|style).*?</\1>", " ", raw)
        raw = re.sub(r"(?i)<br\s*/?>|</(p|div|li|h\d|tr)>", "\n", raw)
        text = re.sub(r"[ \t]{2,}", " ", html.unescape(re.sub(r"<[^>]+>", " ", raw)))
    elif ext in TEXT_EXTS or ext == "":
        if b"\x00" in data[:2048] and not data.startswith((b"\xff\xfe", b"\xfe\xff")):
            raise ValueError(f"That doesn't look like a text file. I can read: {KINDS}.")
        text = _decode(data)
    else:
        raise ValueError(f"I can't read {ext} files yet. I can read: {KINDS}.")
    text = _clean(text)
    if not text:
        raise ValueError("I couldn't find any text in that file.")
    return text


class Library:
    """One person's documents."""

    def __init__(self, path=":memory:", clock=time.time):
        self.db = SafeDB(path)                                   # shared by the server's threads (safe_db.py)
        self.clock = clock
        self.db.execute("CREATE TABLE IF NOT EXISTS docs (id INTEGER PRIMARY KEY, name TEXT, kind TEXT, chars INTEGER, created REAL, text TEXT)")
        try:
            self.db.execute("ALTER TABLE docs ADD COLUMN image BLOB")          # a picture is kept as it was sent, so it can be shown and looked at
            self.db.commit()
        except sqlite3.OperationalError:
            pass
        self.db.commit()

    def add(self, name, data, allow_picture=False):
        name = re.sub(r"[\x00-\x1f/\\]", "", str(name or "document")).strip()[:120] or "document"
        text = extract_text(name, data, allow_picture)
        note = ""
        if len(text) > MAX_CHARS:
            text, note = text[:MAX_CHARS], f"Only the first {MAX_CHARS // 1000:,}K letters were kept."
        if self.db.execute("SELECT COUNT(*) FROM docs").fetchone()[0] >= MAX_DOCS:
            raise ValueError(f"You can keep {MAX_DOCS} documents. Delete one you no longer need first.")
        cur = self.db.execute("INSERT INTO docs (name, kind, chars, created, text, image) VALUES (?, ?, ?, ?, ?, ?)",
                              (name, _ext(name).lstrip(".") or "txt", len(text), self.clock(), text, data if _ext(name) in IMAGE_EXTS else None))
        self.db.commit()
        return dict(self.info(cur.lastrowid), note=note)

    def image(self, doc_id):
        """(bytes, mime type) of a picture that was added, or None."""
        r = self.db.execute("SELECT name, image FROM docs WHERE id = ?", (doc_id,)).fetchone()
        if not r or not r[1]:
            return None
        ext = _ext(r[0])
        return bytes(r[1]), {".jpg": "image/jpeg", ".jpeg": "image/jpeg", ".png": "image/png", ".webp": "image/webp", ".gif": "image/gif", ".bmp": "image/bmp"}.get(ext, "image/png")

    def info(self, doc_id):
        r = self.db.execute("SELECT id, name, kind, chars, created FROM docs WHERE id = ?", (doc_id,)).fetchone()
        return dict(zip(("id", "name", "kind", "chars", "created"), r)) if r else None

    def list(self):
        rows = self.db.execute("SELECT id, name, kind, chars, created FROM docs ORDER BY id DESC").fetchall()
        return [dict(zip(("id", "name", "kind", "chars", "created"), r)) for r in rows]

    def text(self, doc_id):
        r = self.db.execute("SELECT text FROM docs WHERE id = ?", (doc_id,)).fetchone()
        return r[0] if r else None

    def delete(self, doc_id):
        n = self.db.execute("DELETE FROM docs WHERE id = ?", (doc_id,)).rowcount
        self.db.commit()
        return bool(n)

    def clear(self):
        self.db.execute("DELETE FROM docs")
        self.db.commit()


# ------------------------------------------------------------------ choosing the passages for a question
STOP = set("the a an of to in on at for and or is are was were be been do does did what who whom when where which why how can could will would should it its this "
           "that these those with as by from about me my you your i tell give name please document file pdf text say says said".split())
OVERVIEW = re.compile(r"\b(summari[sz]e|summary|overview|main (points?|ideas?)|key (points?|takeaways?|ideas?)|tl;?dr|outline|what is (this|it|the document|the file) about|"
                      r"what does (this|it) (say|cover|contain)|explain (this|the) (document|file|paper|article|text)|give me the gist|in short)\b", re.I)


def chunks(text, size=1200):
    """Pieces of about `size` letters, cut at paragraph or sentence ends."""
    out, cur = [], ""
    for para in re.split(r"\n{2,}|\n(?=\[Page \d+\])", text):
        para = para.strip()
        if not para:
            continue
        while len(para) > size * 1.5:                               # a very long paragraph: cut at a sentence end near `size`
            cut = max(para.rfind(". ", 0, size), para.rfind("\n", 0, size))
            cut = cut + 1 if cut > size // 3 else size
            piece, para = para[:cut].strip(), para[cut:].strip()
            if cur:
                out.append(cur)
                cur = ""
            out.append(piece)
        if len(cur) + len(para) + 2 > size and cur:
            out.append(cur)
            cur = para
        else:
            cur = f"{cur}\n\n{para}" if cur else para
    if cur:
        out.append(cur)
    return out


def _words(text):
    return [w for w in re.findall(r"[a-z0-9']+", text.lower()) if len(w) > 2 and w not in STOP]


def pick(pieces, question, budget):
    """Indexes of the pieces to show, in document order, within `budget` letters."""
    if sum(len(p) for p in pieces) <= budget:
        return list(range(len(pieces)))
    want = set(_words(question))
    chosen, used = [], 0
    if OVERVIEW.search(question) or not want:                       # spread over the whole document: the start, then evenly
        n = max(1, min(len(pieces), budget // max(1, int(sum(len(p) for p in pieces) / len(pieces)))))
        step = len(pieces) / n
        order = sorted({min(len(pieces) - 1, int(i * step)) for i in range(n)} | {0})
    else:
        df = {}
        toks = [_words(p) for p in pieces]
        for t in toks:
            for w in set(t) & want:
                df[w] = df.get(w, 0) + 1
        scores = []
        for i, t in enumerate(toks):
            tf = {}
            for w in t:
                if w in want:
                    tf[w] = tf.get(w, 0) + 1
            s = sum((1 + math.log(c)) * (math.log(1 + len(pieces) / df[w])) for w, c in tf.items())
            scores.append((s, -i))
        order = [i for (s, _), i in sorted(((sc, i) for i, sc in enumerate(scores)), reverse=True) if s > 0]
        if not order:                                               # nothing shares a word with the question: the start of the document
            order = list(range(len(pieces)))
    for i in order:
        if used + len(pieces[i]) > budget and chosen:
            continue
        chosen.append(i)
        used += len(pieces[i])
        if used >= budget:
            break
    return sorted(chosen)


def notes_for(library, doc_ids, question, budget=8000):
    """Notes ([{"title", "text"}]) from the attached documents for this question, within `budget` letters in all."""
    docs = [(i, library.info(i), library.text(i)) for i in doc_ids]
    docs = [(i, info, text) for i, info, text in docs if info and text]
    if not docs:
        return []
    per = max(1500, budget // len(docs))
    notes = []
    for _, info, text in docs:
        if info.get("kind") in IMAGE_KINDS:                    # a picture is not a document: its words go in as what the picture says, with no file name to cite
            words = text.replace("[Text read from a picture]", "").strip()
            if words and not words.startswith(PICTURE_ONLY):
                notes.append({"title": "Words in the picture", "text": words[:per]})
            continue
        pieces = chunks(text)
        idx = pick(pieces, question, per)
        whole = len(idx) == len(pieces)
        if whole:
            notes.append({"title": f"Document: {info['name']}", "text": "\n\n".join(pieces)})
        else:
            for i in idx:
                notes.append({"title": f"Document: {info['name']} (part {i + 1} of {len(pieces)})", "text": pieces[i]})
    return notes


PICTURE_ONLY_TEXT = "(sent a picture)"        # what is saved when only a picture is sent; the page shows just the picture
FILE_ONLY_TEXT = "(sent a file)"              # ...and when only a file is sent (the model then sums it up)


PICTURE_PLAYBOOK = (
    "How to reply, the way a friend who can see the picture would:\n"
    "- Open by showing you got it, in their tone: name what it is and react (\"Caleb's back!\", \"Oof, that error\"). No \"The picture shows\".\n"
    "- Then add something they did not say: what it means, useful context, or your honest take. Two to four short sentences, not a list.\n"
    "- Fit the kind of picture: news or a post: react and add context. An error or code: say what is wrong and how to fix it. Homework, a form "
    "or a bill: help with it. A chart: the main takeaway. A photo: comment on what stands out. A meme: get the joke. A product: what to know before buying.\n"
    "- End with one short question or next step only when it fits naturally.\n"
    "- Do not read out or list its text, do not describe the phone screen around it, and never add sources or file names in brackets.")
TALK_TO_THEM = ("Talk to the user directly as \"you\". If you know their name, it is the name of the person you are talking to: never talk about them as "
                "someone else, and never guess what they will say. Plain, natural sentences.")


def picture_hint(only_picture=False, earlier=False):
    """The instruction for a model that looks at the picture itself: answer like Claude or ChatGPT do with a picture (react, add value, fit the
    kind of picture), not like a text reader quoting a document. `earlier`: the picture came with an earlier message, and the user is now just
    talking: answer what they said first."""
    if earlier:
        return ("[Earlier in this chat the user you are talking with shared the picture you can see here. Now reply to what they just said, like a friend: "
                "react to it first (share their excitement, answer their question, give your take), and bring in the picture only where it helps. Keep it short. "
                "Do not repeat the picture's text and never add sources or file names in brackets. " + TALK_TO_THEM + "]")
    lead = ("[The user you are talking with sent you a picture with no message. You can see it.\n" if only_picture else
            "[The user you are talking with sent you a picture and said something about it. You can see it. Answer what they said first, using what you see.\n")
    return lead + PICTURE_PLAYBOOK + "\n" + TALK_TO_THEM + "]"


def picture_words_hint():
    """For a model that cannot see pictures: it only gets the words read from the picture, and must still talk about it as a picture."""
    return ("[The user you are talking with sent you a picture. You cannot see it, but the notes hold the words in it. Talk about it as their picture, never "
            "as a document or file, and reply like a friend: react first, then add something useful. Do not repeat its text, and never add sources or file "
            "names in brackets. " + TALK_TO_THEM + "]")


def hint(names):
    """The instruction for the model while documents are attached (never saved in the chat)."""
    shown = ", ".join(f'"{n}"' for n in names[:5])
    return (f"[The person attached {shown}. The notes hold the relevant parts of it. Answer from the document, quote a short phrase when it helps "
            "(and say which page or part if it is shown), and say plainly if the document does not contain the answer instead of guessing.]")
