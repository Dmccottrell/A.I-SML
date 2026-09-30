"""
export_doc.py - Turn an AI answer into a printable PDF or an editable Word file.

WHAT THIS FILE DOES
    The model only writes text. The Teacher assistant (skill pack, docs/SKILLS.md) writes it in a simple
    layout that reads fine in the chat and that this file turns into a clean page:

        # Fractions on a Number Line                  <- title
        Grade 3 - Math - Worksheet                    <- plain lines: paragraphs
        Instructions: Write the fraction for each point.
        ## Questions                                  <- a section
        1. What fraction is halfway between 0 and 1?  <- numbered items
        2. ...
        - bullet points work too
        ## Answer key                                 <- always starts on a NEW PAGE
        1. 1/2

    The layout is fixed code, not the model's job, so every worksheet looks the same and prints well:
    big readable fonts, "Name: ____  Date: ____" on student pages, writing space under each question,
    and the answer key on its own page (print page 1 for the class, keep the key).

    Everything runs offline on your computer. Needs:  pip install fpdf2 python-docx

Usage:
    python export_doc.py answer.txt --out worksheet.pdf          (or .docx, .md, .txt)
    In chat (generate.py --chat):  save worksheet.pdf            (saves the last answer into documents/)
"""
import argparse
import os
import re
from dataclasses import dataclass, field

DOCS_DIR = "documents"
FORMATS = (".pdf", ".docx", ".md", ".txt")


@dataclass
class Section:
    heading: str = ""
    blocks: list = field(default_factory=list)     # ("para", text) | ("numbered", [items]) | ("bullets", [items])

    @property
    def is_key(self):
        return bool(re.search(r"\banswer key\b|\banswers\b", self.heading, re.I))

    @property
    def is_questions(self):
        return bool(re.search(r"\bquestions?\b|\bproblems?\b|\bpractice\b|\bexercises?\b", self.heading, re.I))


@dataclass
class Doc:
    title: str = ""
    sections: list = field(default_factory=list)

    @property
    def for_students(self):
        """A page students write on (worksheets, reading questions): gets the Name/Date line."""
        return any(s.is_questions for s in self.sections)

    def items(self, which):
        """Numbered items of the question sections ("questions") or the answer key ("key")."""
        want = (lambda s: s.is_key) if which == "key" else (lambda s: s.is_questions and not s.is_key)
        return [it for s in self.sections if want(s) for kind, v in s.blocks if kind == "numbered" for it in v]


NUMBERED = re.compile(r"^\s*(\d+)[.)]\s+(.*)$")
BULLET = re.compile(r"^\s*[-*•]\s+(.*)$")


def parse(text):
    """Doc from the model's text. Lines that don't match anything special become paragraphs."""
    doc = Doc()
    cur = Section()
    doc.sections.append(cur)
    for raw in text.replace("\r\n", "\n").split("\n"):
        line = raw.strip()
        if not line:
            continue
        line = re.sub(r"\*\*(.+?)\*\*", r"\1", line)                  # **bold** -> plain
        if line.startswith("# ") and not doc.title:
            doc.title = line[2:].strip()
            continue
        if line.startswith("#"):
            cur = Section(line.lstrip("#").strip())
            doc.sections.append(cur)
            continue
        m, b = NUMBERED.match(line), BULLET.match(line)
        if m or b:
            kind, item = ("numbered", m.group(2)) if m else ("bullets", b.group(1))
            if cur.blocks and cur.blocks[-1][0] == kind:
                cur.blocks[-1][1].append(item)
            else:
                cur.blocks.append((kind, [item]))
        elif cur.blocks and cur.blocks[-1][0] in ("numbered", "bullets") and raw.startswith(("   ", "\t")):
            cur.blocks[-1][1][-1] += " " + line                       # an item that continues on the next line
        else:
            cur.blocks.append(("para", line))
    doc.sections = [s for s in doc.sections if s.heading or s.blocks]
    if not doc.title:
        doc.title = "Document"
    return doc


# ------------------------------------------------------------------ writers
_ASCII = {"‘": "'", "’": "'", "“": '"', "”": '"', "–": "-", "—": "-",
          "•": "-", "…": "...", " ": " ", "−": "-", "→": "->", "≤": "<=",
          "≥": ">="}


def _latin1(s):
    """The PDF's built-in fonts only know Latin-1: swap curly quotes and dashes for plain ones."""
    s = "".join(_ASCII.get(c, c) for c in s)
    return s.encode("latin-1", "replace").decode("latin-1")


def to_pdf(doc, path):
    from fpdf import FPDF
    pdf = FPDF(format="Letter")
    pdf.set_margins(20, 18, 20)
    pdf.set_auto_page_break(True, 18)
    pdf.add_page()
    w = pdf.w - pdf.l_margin - pdf.r_margin

    def text(s, size=14, style="", h=None, indent=0):
        pdf.set_font("Helvetica", style, size)
        pdf.set_x(pdf.l_margin + indent)
        pdf.multi_cell(w - indent, h or size * 0.55, _latin1(s))

    text(doc.title, 22, "B", 11)
    if doc.for_students:
        pdf.ln(2)
        text("Name: ______________________________    Date: ______________", 13)
    pdf.ln(4)
    for sec in doc.sections:
        if sec.is_key:
            pdf.add_page()
            text(doc.title, 11, "I")                                  # which worksheet this key belongs to
        if sec.heading:
            pdf.ln(2)
            text(sec.heading, 17, "B", 9)
            pdf.ln(1)
        for kind, v in sec.blocks:
            if kind == "para":
                text(v)
                pdf.ln(1.5)
            else:
                for i, item in enumerate(v, 1):
                    text((f"{i}. " if kind == "numbered" else "-  ") + item, indent=2)
                    if kind == "numbered" and sec.is_questions and not sec.is_key:
                        pdf.ln(5 if "____" in item else 16)          # room to write the answer
                    else:
                        pdf.ln(1.5)
    pdf.output(path)


def to_docx(doc, path):
    from docx import Document
    from docx.enum.text import WD_BREAK
    from docx.shared import Pt
    d = Document()
    d.styles["Normal"].font.size = Pt(14)
    d.add_heading(doc.title, level=0)
    if doc.for_students:
        d.add_paragraph("Name: ______________________________    Date: ______________")
    for sec in doc.sections:
        if sec.is_key:
            d.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        if sec.heading:
            d.add_heading(sec.heading, level=1)
        for kind, v in sec.blocks:
            if kind == "para":
                d.add_paragraph(v)
            else:
                for i, item in enumerate(v, 1):
                    d.add_paragraph((f"{i}. " if kind == "numbered" else "• ") + item)
                    if kind == "numbered" and sec.is_questions and not sec.is_key and "____" not in item:
                        d.add_paragraph("")                           # room to write the answer
    d.save(path)


def to_markdown(doc):
    out = [f"# {doc.title}", ""]
    if doc.for_students:
        out += ["Name: ____________  Date: ________", ""]
    for sec in doc.sections:
        if sec.heading:
            out += [f"## {sec.heading}", ""]
        for kind, v in sec.blocks:
            if kind == "para":
                out += [v, ""]
            else:
                out += [(f"{i}. " if kind == "numbered" else "- ") + it for i, it in enumerate(v, 1)] + [""]
    return "\n".join(out).rstrip() + "\n"


def export(text, path):
    """Write `text` (the model's answer) to `path`; the format comes from the extension. Returns the path."""
    ext = os.path.splitext(path)[1].lower()
    if ext not in FORMATS:
        raise ValueError(f"can't save as {ext or '(no extension)'}: use {', '.join(FORMATS)}")
    os.makedirs(os.path.dirname(path) or ".", exist_ok=True)
    text = re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f\x7f]", "", text)       # invisible control characters
    doc = parse(text)
    try:
        if ext == ".pdf":
            to_pdf(doc, path)
        elif ext == ".docx":
            to_docx(doc, path)
        else:
            with open(path, "w", encoding="utf-8") as f:
                f.write(to_markdown(doc) if ext == ".md" else text.strip() + "\n")
    except ImportError as e:
        raise SystemExit(f"saving {ext} files needs an extra package: pip install fpdf2 python-docx ({e})") from None
    return path


def safe_name(name, default="document.pdf"):
    """A file name inside documents/ (no folders, no odd characters); .pdf if no known extension."""
    base = os.path.basename(name.strip()) or default
    base = re.sub(r"[^A-Za-z0-9._ -]", "", base).strip(" .") or default
    if os.path.splitext(base)[1].lower() not in FORMATS:
        base += ".pdf"
    return os.path.join(DOCS_DIR, base)


def main():
    p = argparse.ArgumentParser(description="Save an AI answer as a PDF, Word, Markdown or text file")
    p.add_argument("input", help="a text file with the answer")
    p.add_argument("--out", required=True, help="e.g. worksheet.pdf or worksheet.docx")
    a = p.parse_args()
    with open(a.input, encoding="utf-8") as f:
        print("saved", export(f.read(), a.out))


if __name__ == "__main__":
    main()
